from __future__ import annotations

import math
import os
import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT_DOCX = ROOT / "初稿_完成版.docx"
IMG_DIR = ROOT / "images"
IMG_DIR.mkdir(exist_ok=True)

TITLE = "面向普通人群的血糖健康管理助手 APP 的设计与实现"
HEADER_TEXT = "大学本科毕业论文（设计）"

PRIMARY = (54, 135, 124)
PRIMARY_DARK = (32, 86, 79)
GREEN = (62, 151, 107)
YELLOW = (206, 151, 57)
RED = (188, 76, 72)
LAVENDER = (108, 94, 170)
TEXT = (38, 48, 53)
MUTED = (104, 116, 124)
LINE = (213, 224, 224)
SURFACE = (255, 255, 255)
BG = (245, 249, 248)


def pick_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = []
    if bold:
        candidates += [
            r"C:\Windows\Fonts\simhei.ttf",
            r"C:\Windows\Fonts\msyhbd.ttc",
        ]
    candidates += [
        r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\simsun.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
    ]
    for path in candidates:
        if os.path.exists(path):
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


FONT_16 = pick_font(16)
FONT_18 = pick_font(18)
FONT_20 = pick_font(20)
FONT_22B = pick_font(22, True)
FONT_24B = pick_font(24, True)
FONT_28B = pick_font(28, True)
FONT_32B = pick_font(32, True)
FONT_36B = pick_font(36, True)


def draw_round(draw: ImageDraw.ImageDraw, box, radius=18, fill=SURFACE, outline=LINE, width=2):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def text_size(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont):
    bbox = draw.textbbox((0, 0), text, font=font)
    return bbox[2] - bbox[0], bbox[3] - bbox[1]


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font, max_width: int):
    lines = []
    current = ""
    for ch in text:
        candidate = current + ch
        if text_size(draw, candidate, font)[0] <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = ch
    if current:
        lines.append(current)
    return lines


def draw_text_box(draw, box, text, font, fill=TEXT, align="center", line_gap=6):
    x1, y1, x2, y2 = box
    max_width = x2 - x1 - 24
    lines = []
    for part in text.split("\n"):
        lines.extend(wrap_text(draw, part, font, max_width) or [""])
    total_h = sum(text_size(draw, line, font)[1] for line in lines) + line_gap * (len(lines) - 1)
    y = y1 + (y2 - y1 - total_h) / 2
    for line in lines:
        w, h = text_size(draw, line, font)
        if align == "left":
            x = x1 + 16
        elif align == "right":
            x = x2 - w - 16
        else:
            x = x1 + (x2 - x1 - w) / 2
        draw.text((x, y), line, font=font, fill=fill)
        y += h + line_gap


def arrow(draw, start, end, fill=PRIMARY_DARK, width=3):
    draw.line([start, end], fill=fill, width=width)
    x1, y1 = start
    x2, y2 = end
    ang = math.atan2(y2 - y1, x2 - x1)
    size = 10
    pts = [
        (x2, y2),
        (x2 - size * math.cos(ang - math.pi / 6), y2 - size * math.sin(ang - math.pi / 6)),
        (x2 - size * math.cos(ang + math.pi / 6), y2 - size * math.sin(ang + math.pi / 6)),
    ]
    draw.polygon(pts, fill=fill)


def save_use_case():
    img = Image.new("RGB", (1500, 900), BG)
    d = ImageDraw.Draw(img)
    d.text((50, 36), "系统用例图", font=FONT_36B, fill=PRIMARY_DARK)
    draw_round(d, (330, 115, 1200, 800), 30, (255, 255, 255), LINE, 3)
    d.text((642, 140), "血糖健康管理助手 APP", font=FONT_28B, fill=PRIMARY_DARK)
    actor_x, actor_y = 150, 420
    d.ellipse((actor_x - 28, actor_y - 100, actor_x + 28, actor_y - 44), outline=TEXT, width=4)
    d.line((actor_x, actor_y - 44, actor_x, actor_y + 55), fill=TEXT, width=4)
    d.line((actor_x - 55, actor_y - 10, actor_x + 55, actor_y - 10), fill=TEXT, width=4)
    d.line((actor_x, actor_y + 55, actor_x - 48, actor_y + 130), fill=TEXT, width=4)
    d.line((actor_x, actor_y + 55, actor_x + 48, actor_y + 130), fill=TEXT, width=4)
    d.text((105, 570), "普通用户", font=FONT_24B, fill=TEXT)

    clock_x, clock_y = 1340, 410
    d.ellipse((clock_x - 55, clock_y - 55, clock_x + 55, clock_y + 55), outline=PRIMARY_DARK, width=4)
    d.line((clock_x, clock_y, clock_x, clock_y - 35), fill=PRIMARY_DARK, width=4)
    d.line((clock_x, clock_y, clock_x + 30, clock_y + 20), fill=PRIMARY_DARK, width=4)
    d.text((1260, 500), "时间触发器\n提醒服务", font=FONT_22B, fill=PRIMARY_DARK)

    cases = [
        (430, 220, "注册/登录"),
        (690, 220, "维护个人资料"),
        (950, 220, "设置每日提醒"),
        (430, 390, "记录血糖"),
        (690, 390, "记录饮食"),
        (950, 390, "识别食物图片"),
        (430, 560, "记录运动"),
        (690, 560, "记录主观状态"),
        (950, 560, "查看分析建议"),
        (690, 710, "接收提醒并回看数据"),
    ]
    for x, y, label in cases:
        d.ellipse((x - 110, y - 44, x + 110, y + 44), fill=(238, 248, 246), outline=PRIMARY, width=3)
        draw_text_box(d, (x - 104, y - 38, x + 104, y + 38), label, FONT_20, PRIMARY_DARK)
        if x < 800:
            arrow(d, (230, 460), (x - 115, y), MUTED, 2)
    for target in [(950, 220), (690, 710)]:
        arrow(d, (1280, 440), (target[0] + 116, target[1]), LAVENDER, 2)
    img.save(IMG_DIR / "fig_use_case.png", quality=95)


def save_architecture():
    img = Image.new("RGB", (1500, 900), BG)
    d = ImageDraw.Draw(img)
    d.text((50, 36), "系统总体架构图", font=FONT_36B, fill=PRIMARY_DARK)
    layers = [
        (90, 145, 1410, 265, "表示层：Flutter Material 页面", "登录、首页、血糖、饮食、运动、分析、个人设置"),
        (90, 310, 1410, 430, "状态与业务层：Provider + Repository + Service", "HealthProvider 管理页面状态；HealthRepository 封装 Supabase 数据访问；AnalysisService 执行本地规则计算"),
        (90, 475, 1410, 625, "BaaS 后端层：Supabase", "Auth 认证、Postgres 数据库、Storage 私有图片桶、Edge Function 后端中转"),
        (90, 670, 1410, 800, "外部服务层：百度菜品识别 API", "Edge Function 获取 access_token 后调用菜品识别接口，客户端不接触百度密钥"),
    ]
    colors = [(232, 247, 244), (242, 241, 251), (232, 241, 251), (255, 247, 232)]
    for idx, (x1, y1, x2, y2, title, desc) in enumerate(layers):
        draw_round(d, (x1, y1, x2, y2), 22, colors[idx], LINE, 3)
        d.text((x1 + 30, y1 + 20), title, font=FONT_28B, fill=PRIMARY_DARK)
        d.text((x1 + 30, y1 + 68), desc, font=FONT_22B, fill=TEXT)
    for y in [265, 430, 625]:
        arrow(d, (750, y + 10), (750, y + 42), PRIMARY_DARK, 4)
    d.text((1020, 500), "supabase_flutter SDK\nHTTPS 调用", font=FONT_22B, fill=PRIMARY_DARK)
    img.save(IMG_DIR / "fig_architecture.png", quality=95)


def save_er():
    img = Image.new("RGB", (1700, 1100), BG)
    d = ImageDraw.Draw(img)
    d.text((50, 36), "简化 ER 图", font=FONT_36B, fill=PRIMARY_DARK)
    entities = {
        "user_profile\nPK id": (705, 105, 995, 205),
        "user_body_metrics\nFK user_id": (35, 360, 285, 465),
        "blood_glucose_logs\nFK user_id": (305, 360, 575, 465),
        "meals\nFK user_id": (600, 360, 850, 465),
        "exercise_logs\nFK user_id": (875, 360, 1145, 465),
        "wellness_status\nFK user_id\n可关联餐食/运动": (1165, 345, 1415, 490),
        "reminder_settings\nFK user_id": (1435, 360, 1685, 465),
        "meal_items\nFK meal_id": (600, 635, 850, 740),
        "exercise_catalog\n运动 MET 目录": (875, 635, 1145, 740),
        "food_calorie_catalog\n食物热量目录\n非强制外键": (600, 850, 850, 985),
    }
    for name, box in entities.items():
        draw_round(d, box, 18, SURFACE, PRIMARY if "catalog" not in name else LAVENDER, 3)
        draw_text_box(d, box, name, FONT_20, TEXT)
    center = {
        name: ((box[0] + box[2]) // 2, (box[1] + box[3]) // 2)
        for name, box in entities.items()
    }
    def port(name, side):
        x1, y1, x2, y2 = entities[name]
        if side == "top":
            return ((x1 + x2) // 2, y1)
        if side == "bottom":
            return ((x1 + x2) // 2, y2)
        if side == "left":
            return (x1, (y1 + y2) // 2)
        return (x2, (y1 + y2) // 2)

    def draw_poly(points, label=None, label_at=None):
        d.line(points, fill=PRIMARY_DARK, width=3)
        if label:
            x, y = label_at or points[len(points) // 2]
            draw_round(d, (x - 54, y - 18, x + 54, y + 18), 10, (255, 255, 255), LINE, 1)
            draw_text_box(d, (x - 50, y - 16, x + 50, y + 16), label, FONT_16, PRIMARY_DARK)

    # User-owned rows are connected through a top bus line to keep relationship
    # labels readable and avoid crossing entity text.
    user_bottom = port("user_profile\nPK id", "bottom")
    bus_y = 285
    child_names = [
        "user_body_metrics\nFK user_id",
        "blood_glucose_logs\nFK user_id",
        "meals\nFK user_id",
        "exercise_logs\nFK user_id",
        "wellness_status\nFK user_id\n可关联餐食/运动",
        "reminder_settings\nFK user_id",
    ]
    xs = [port(name, "top")[0] for name in child_names]
    draw_poly([user_bottom, (user_bottom[0], bus_y), (min(xs), bus_y), (max(xs), bus_y)])
    for name in child_names:
        top = port(name, "top")
        draw_poly([(top[0], bus_y), top], "1:N", (top[0], (bus_y + top[1]) // 2))

    draw_poly([port("meals\nFK user_id", "bottom"), port("meal_items\nFK meal_id", "top")], "1:N")
    draw_poly([port("exercise_catalog\n运动 MET 目录", "top"), port("exercise_logs\nFK user_id", "bottom")], "1:N")
    draw_poly([
        port("food_calorie_catalog\n食物热量目录\n非强制外键", "top"),
        (port("food_calorie_catalog\n食物热量目录\n非强制外键", "top")[0], 770),
        (port("meal_items\nFK meal_id", "bottom")[0], 770),
        port("meal_items\nFK meal_id", "bottom"),
    ], "查询", (900, 770))
    img.save(IMG_DIR / "fig_er.png", quality=95)


def save_sequence():
    img = Image.new("RGB", (1700, 1000), BG)
    d = ImageDraw.Draw(img)
    d.text((50, 36), "饮食图片识别时序图", font=FONT_36B, fill=PRIMARY_DARK)
    actors = [
        ("用户", 170),
        ("Flutter 客户端", 470),
        ("Supabase Storage", 800),
        ("Edge Function", 1130),
        ("百度菜品识别 API", 1460),
    ]
    for name, x in actors:
        draw_round(d, (x - 115, 120, x + 115, 185), 16, SURFACE, PRIMARY, 3)
        draw_text_box(d, (x - 110, 124, x + 110, 181), name, FONT_20, PRIMARY_DARK)
        d.line((x, 185, x, 900), fill=(180, 196, 196), width=2)
    steps = [
        (170, 470, 250, "拍照/选择图片"),
        (470, 800, 330, "上传到私有图片桶"),
        (470, 1130, 420, "调用 recognize-meal"),
        (1130, 800, 510, "service role 下载用户图片"),
        (1130, 1460, 600, "获取 access_token"),
        (1130, 1460, 690, "提交图片 Base64"),
        (1460, 1130, 780, "返回 name、calorie、probability"),
        (1130, 470, 860, "返回菜名、热量和置信度"),
    ]
    for x1, x2, y, label in steps:
        arrow(d, (x1 + (35 if x2 > x1 else -35), y), (x2 - (35 if x2 > x1 else -35), y), PRIMARY_DARK, 3)
        d.text((min(x1, x2) + 25, y - 34), label, font=FONT_18, fill=TEXT)
    d.text((1020, 925), "当前实现按请求获取 access_token；后续可加入缓存降低调用延迟。", font=FONT_18, fill=MUTED)
    img.save(IMG_DIR / "fig_sequence.png", quality=95)


def save_data_flow():
    img = Image.new("RGB", (1500, 850), BG)
    d = ImageDraw.Draw(img)
    d.text((50, 36), "核心数据流图", font=FONT_36B, fill=PRIMARY_DARK)
    boxes = [
        (80, 170, 320, 290, "用户输入\n血糖/饮食/运动/状态"),
        (430, 170, 710, 290, "HealthProvider\n刷新页面状态"),
        (820, 170, 1120, 290, "HealthRepository\nSupabase CRUD"),
        (1210, 170, 1440, 290, "Postgres\n用户私有数据"),
        (430, 520, 710, 640, "AnalysisService\n纯业务规则计算"),
        (820, 520, 1120, 640, "首页/分析页\n红黄绿灯与建议"),
    ]
    for box in boxes:
        draw_round(d, box[:4], 20, SURFACE, PRIMARY, 3)
        draw_text_box(d, box[:4], box[4], FONT_22B, TEXT)
    arrows = [
        ((320, 230), (430, 230)),
        ((710, 230), (820, 230)),
        ((1120, 230), (1210, 230)),
        ((960, 295), (620, 520)),
        ((710, 580), (820, 580)),
        ((575, 520), (575, 295)),
    ]
    for a, b in arrows:
        arrow(d, a, b, PRIMARY_DARK, 4)
    d.text((120, 725), "数据先通过后端安全策略进入用户行级空间，再由本地规则生成可解释反馈；客户端不持有服务端密钥。", font=FONT_22B, fill=PRIMARY_DARK)
    img.save(IMG_DIR / "fig_data_flow.png", quality=95)


def phone_frame(draw, x, y, title, cards, accent=PRIMARY):
    draw_round(draw, (x, y, x + 310, y + 620), 34, (24, 35, 39), (24, 35, 39), 2)
    draw_round(draw, (x + 14, y + 14, x + 296, y + 606), 28, BG, (24, 35, 39), 1)
    draw.rounded_rectangle((x + 110, y + 28, x + 200, y + 40), radius=6, fill=(24, 35, 39))
    draw.text((x + 32, y + 70), title, font=FONT_24B, fill=PRIMARY_DARK)
    yy = y + 118
    for card in cards:
        h = card.get("h", 92)
        draw_round(draw, (x + 30, yy, x + 280, yy + h), 14, SURFACE, (226, 236, 235), 2)
        draw.text((x + 48, yy + 18), card["title"], font=FONT_18 if len(card["title"]) > 10 else FONT_20, fill=TEXT)
        draw.text((x + 48, yy + 48), card["body"], font=FONT_16, fill=card.get("color", MUTED))
        if "dot" in card:
            draw.ellipse((x + 238, yy + 28, x + 258, yy + 48), fill=card["dot"])
        yy += h + 18
    draw.rounded_rectangle((x + 70, y + 570, x + 240, y + 590), radius=10, fill=accent)


def save_screens():
    img = Image.new("RGB", (1800, 1420), (239, 246, 245))
    d = ImageDraw.Draw(img)
    d.text((60, 44), "系统主要页面实现效果（演示数据，已脱敏）", font=FONT_36B, fill=PRIMARY_DARK)
    screens = [
        ("登录", [
            {"title": "稳啦", "body": "血糖健康管理助手", "h": 105, "color": PRIMARY_DARK},
            {"title": "邮箱", "body": "demo@example.com", "h": 82},
            {"title": "密码", "body": "••••••••", "h": 82},
            {"title": "进入稳啦", "body": "Supabase Auth", "h": 82, "color": PRIMARY},
        ], PRIMARY),
        ("首页", [
            {"title": "今天状态稳吗？", "body": "最新血糖 6.4 mmol/L", "h": 100, "color": PRIMARY_DARK},
            {"title": "红黄绿灯食物", "body": "燕麦牛奶 绿灯", "h": 90, "dot": GREEN},
            {"title": "饭后轻运动", "body": "餐后慢走 10-15 分钟", "h": 90},
            {"title": "血糖波动", "body": "3.9-10.0 参考区间", "h": 120},
        ], GREEN),
        ("血糖", [
            {"title": "6.4 mmol/L", "body": "平稳", "h": 110, "color": PRIMARY},
            {"title": "记录时段", "body": "早餐后 / 睡前 / 随机", "h": 90},
            {"title": "来源", "body": "手动 / CGM 预留", "h": 90},
            {"title": "历史记录", "body": "04-30 09:30 早餐后", "h": 90},
        ], PRIMARY),
        ("饮食", [
            {"title": "拍一下这餐", "body": "拍照 / 相册 / 手动加", "h": 95, "color": PRIMARY_DARK},
            {"title": "确认食物", "body": "燕麦粥 250g", "h": 90},
            {"title": "热量", "body": "68 kcal/100g, 170 kcal", "h": 90},
            {"title": "最近吃过", "body": "米饭 150g 174 kcal", "h": 92},
        ], YELLOW),
        ("分析", [
            {"title": "手动记录", "body": "今天整体比较稳", "h": 105, "color": PRIMARY_DARK},
            {"title": "可能有关", "body": "餐后波动与主食量有关", "h": 92},
            {"title": "下次试试", "body": "减少甜饮并饭后慢走", "h": 92},
            {"title": "CGM", "body": "连续血糖作为扩展", "h": 92},
        ], LAVENDER),
    ]
    xs = [80, 420, 760, 1100, 1440]
    for i, (title, cards, accent) in enumerate(screens):
        phone_frame(d, xs[i], 145, title, cards, accent)
    d.text((80, 805), "说明：页面图用于论文展示，不包含真实姓名、头像、设备信息或真实健康数据。", font=FONT_22B, fill=RED)
    d.text((80, 860), "底部导航结构：首页、血糖、饮食、运动、分析；个人设置和状态记录由首页入口进入。", font=FONT_22B, fill=PRIMARY_DARK)
    img.save(IMG_DIR / "fig_screens.png", quality=95)


def generate_images():
    save_use_case()
    save_architecture()
    save_er()
    save_sequence()
    save_data_flow()
    save_screens()


body_text_parts: list[str] = []


def set_run_font(run, east="宋体", west="Times New Roman", size=None, bold=None, color=None):
    run.font.name = west
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)
    run._element.rPr.rFonts.set(qn("w:ascii"), west)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), west)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def set_para_format(p, first_indent=True, align=None, line=20, before=0, after=0):
    fmt = p.paragraph_format
    fmt.first_line_indent = Pt(21) if first_indent else Pt(0)
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(line)
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    if align is not None:
        p.alignment = align


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(text)
    run._r.append(fld_end)
    set_run_font(run, size=9)


def configure_document(doc: Document):
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(3.0)
    sec.right_margin = Cm(2.5)
    sec.top_margin = Cm(3.0)
    sec.bottom_margin = Cm(2.5)
    sec.header_distance = Cm(2.3)
    sec.footer_distance = Cm(1.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(10.5)
    pf = normal.paragraph_format
    pf.first_line_indent = Pt(21)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(20)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)

    h1 = doc.styles["Heading 1"]
    h1.font.name = "Times New Roman"
    h1._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    h1._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    h1._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0, 0, 0)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.first_line_indent = Pt(0)
    h1.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    h1.paragraph_format.line_spacing = Pt(20)
    h1.paragraph_format.space_before = Pt(0)
    h1.paragraph_format.space_after = Pt(12)

    h2 = doc.styles["Heading 2"]
    h2.font.name = "Times New Roman"
    h2._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    h2._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    h2._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    h2.font.size = Pt(14)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(0, 0, 0)
    h2.paragraph_format.first_line_indent = Pt(0)
    h2.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    h2.paragraph_format.line_spacing = Pt(20)
    h2.paragraph_format.space_before = Pt(8)
    h2.paragraph_format.space_after = Pt(4)

    h3 = doc.styles["Heading 3"]
    h3.font.name = "Times New Roman"
    h3._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    h3._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    h3._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    h3.font.size = Pt(10.5)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor(0, 0, 0)
    h3.paragraph_format.first_line_indent = Pt(0)
    h3.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    h3.paragraph_format.line_spacing = Pt(20)
    h3.paragraph_format.space_before = Pt(6)
    h3.paragraph_format.space_after = Pt(3)

    header = sec.header
    header_p = header.paragraphs[0]
    header_p.text = HEADER_TEXT
    header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para_format(header_p, first_indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, line=14)
    for run in header_p.runs:
        set_run_font(run, size=9)

    footer_p = sec.footer.paragraphs[0]
    footer_p.text = ""
    add_page_field(footer_p)


def heading(doc, text, level=1, page_break=False):
    if page_break and len(doc.paragraphs) > 0:
        doc.add_page_break()
    p = doc.add_paragraph(text, style=f"Heading {level}")
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in p.runs:
        if level == 1:
            set_run_font(run, east="黑体", size=16, bold=True)
        elif level == 2:
            set_run_font(run, east="黑体", size=14, bold=True)
        else:
            set_run_font(run, east="宋体", size=10.5, bold=True)
    return p


def para(doc, text, collect=True):
    p = doc.add_paragraph()
    set_para_format(p, first_indent=True)
    run = p.add_run(text)
    set_run_font(run, size=10.5)
    if collect:
        body_text_parts.append(text)
    return p


def noindent_para(doc, text, align=None, size=10.5, bold=False):
    p = doc.add_paragraph()
    set_para_format(p, first_indent=False, align=align)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def caption(doc, text):
    p = doc.add_paragraph()
    set_para_format(p, first_indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, line=18, before=3, after=4)
    run = p.add_run(text)
    set_run_font(run, east="黑体", size=9, bold=True)
    return p


def add_picture(doc, img_name, cap, width=5.9):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para_format(p, first_indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = p.add_run()
    run.add_picture(str(IMG_DIR / img_name), width=Inches(width))
    drawings = run._element.xpath(".//wp:docPr")
    for doc_pr in drawings:
        doc_pr.set("descr", cap)
        doc_pr.set("title", cap)
    caption(doc, cap)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_pr.append(parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls("w"), fill)))


def format_cell(cell, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=9):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for p in cell.paragraphs:
        p.alignment = align
        set_para_format(p, first_indent=False, align=align, line=15, before=1, after=1)
        for run in p.runs:
            set_run_font(run, size=font_size, bold=bold)


def table(doc, cap, headers, rows, widths=None):
    caption(doc, cap)
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"
    tbl.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    hdr_cells = tbl.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        shade_cell(hdr_cells[i], "DCEFEB")
        format_cell(hdr_cells[i], bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row in rows:
        cells = tbl.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            align = WD_ALIGN_PARAGRAPH.LEFT if len(str(value)) > 18 else WD_ALIGN_PARAGRAPH.CENTER
            format_cell(cells[i], align=align)
    if widths:
        for row in tbl.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    p = doc.add_paragraph()
    set_para_format(p, first_indent=False, line=10)
    return tbl


def code_block(doc, text):
    for line in text.strip("\n").split("\n"):
        p = doc.add_paragraph()
        set_para_format(p, first_indent=False, line=14, before=0, after=0)
        run = p.add_run(line)
        set_run_font(run, east="等线", west="Consolas", size=8.5)
        p._p.get_or_add_pPr().append(parse_xml(r'<w:shd {} w:fill="F3F6F6"/>'.format(nsdecls("w"))))


def formula(doc, text):
    p = doc.add_paragraph()
    set_para_format(p, first_indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, line=20, before=4, after=4)
    run = p.add_run(text)
    set_run_font(run, east="宋体", west="Times New Roman", size=10.5, bold=True)


def add_toc(doc):
    heading(doc, "目  录", 1)
    entries = [
        ("摘  要", "Ⅰ", 0),
        ("ABSTRACT", "Ⅱ", 0),
        ("第一章 绪论", "1", 0),
        ("1.1 研究背景", "1", 1),
        ("1.2 国内外研究现状", "2", 1),
        ("1.3 研究意义", "3", 1),
        ("1.4 论文主要工作", "4", 1),
        ("第二章 需求分析", "5", 0),
        ("2.1 用户需求与系统边界", "5", 1),
        ("2.2 功能需求分析", "6", 1),
        ("2.3 非功能与安全需求", "7", 1),
        ("第三章 系统总体设计", "9", 0),
        ("3.1 技术路线与架构选择", "9", 1),
        ("3.2 系统架构设计", "10", 1),
        ("3.3 核心数据流设计", "11", 1),
        ("第四章 数据库与安全设计", "13", 0),
        ("4.1 数据模型设计", "13", 1),
        ("4.2 后端安全与 RLS 策略", "15", 1),
        ("4.3 饮食图片存储与密钥管理", "17", 1),
        ("第五章 功能模块实现", "18", 0),
        ("5.1 登录注册与个人档案", "18", 1),
        ("5.2 血糖、饮食、运动和状态记录", "19", 1),
        ("5.3 饮食识别与热量估算", "21", 1),
        ("5.4 分析建议与红黄绿灯规则", "23", 1),
        ("5.5 提醒与界面实现", "25", 1),
        ("第六章 测试与结果分析", "27", 0),
        ("6.1 测试环境与测试方法", "27", 1),
        ("6.2 功能、算法与安全测试", "28", 1),
        ("6.3 界面效果与结果分析", "30", 1),
        ("第七章 结论与展望", "32", 0),
        ("参考文献", "34", 0),
        ("附  录", "36", 0),
        ("致  谢", "38", 0),
    ]
    for title, page, level in entries:
        p = doc.add_paragraph()
        set_para_format(p, first_indent=False, line=20)
        p.paragraph_format.left_indent = Pt(21 * level)
        run = p.add_run(f"{title}{'.' * max(4, 38 - len(title) - level * 2)}{page}")
        set_run_font(run, size=10.5)


def add_abstracts(doc):
    heading(doc, "摘  要", 1, page_break=True)
    paras = [
        "随着居民生活方式变化和慢性病管理需求增加，普通人群对血糖、饮食、运动和身体状态之间关系的关注不断提高。传统血糖记录工具往往侧重单一数值保存，难以帮助用户理解日常行为与血糖波动之间的联系。本文围绕“面向普通人群的血糖健康管理助手 APP 的设计与实现”开展应用开发研究，设计并实现一套以 Flutter 为客户端、Supabase 为后端服务的跨端健康管理系统。",
        "系统面向非专业用户，提供登录注册、血糖记录、饮食记录、食物图片识别、运动记录、主观状态记录、提醒设置和数据分析等功能。客户端采用 Provider 管理界面状态，通过 supabase_flutter SDK 与 Supabase Auth、Postgres、Storage 和 Edge Function 通信；后端使用行级安全策略隔离用户数据，并通过 Edge Function 中转百度菜品识别 API，避免在客户端暴露第三方密钥。系统在分析层采用可解释规则，将餐后 30—180 分钟窗口内的血糖记录、餐食、运动和主观状态进行关联，生成红黄绿灯食物提示和饭后运动建议。",
        "测试结果表明，项目通过 flutter analyze 静态检查，flutter test 中 22 个测试全部通过，覆盖热量计算、MET 运动消耗、红黄绿灯分类、异常提示、个人资料刷新和页面防连击等关键逻辑。本文实现的系统不是医疗诊断工具，不提供疾病诊断或治疗建议，而是为普通用户提供日常记录、趋势观察和生活方式辅助。该设计能够为毕业设计场景下的移动健康应用开发提供一个较完整的工程实现案例。",
    ]
    for item in paras:
        para(doc, item, collect=False)
    noindent_para(doc, "关键词：血糖管理；Flutter；Supabase；健康数据分析", bold=True)

    heading(doc, "ABSTRACT", 1, page_break=True)
    english = [
        "With the increasing demand for lifestyle-based chronic disease management, ordinary users need tools that can connect blood glucose records with meals, exercise, and subjective wellness status. Many existing record tools focus mainly on storing isolated values, while they provide limited support for explaining why a user may feel tired, hungry, or unstable after certain daily behaviors. This thesis designs and implements a blood glucose health management assistant application for general users. The system uses Flutter as the cross-platform client and Supabase as the Backend as a Service platform.",
        "The application provides authentication, blood glucose logging, meal logging, food image recognition, exercise logging, wellness status logging, reminders, and data analysis. On the client side, Provider is used to manage UI state, and the supabase_flutter SDK is used to communicate with Supabase Auth, Postgres, Storage, and Edge Functions. On the backend side, Row Level Security policies isolate user data, and an Edge Function acts as a secure relay for Baidu dish recognition API calls so that third-party secrets are not exposed to the Flutter client. The analysis module uses explainable rules to associate meals with blood glucose records in the 30 to 180 minute post-meal window and generates traffic-light food signals and post-meal exercise suggestions.",
        "The project passed flutter analyze, and all 22 flutter test cases passed. The tests cover calorie calculation, MET-based exercise energy estimation, traffic-light food classification, exception messages, profile refresh, and button anti-repeat behavior. The system is not intended for medical diagnosis, treatment decisions, or replacement of professional medical advice. Its value lies in helping ordinary users keep daily records, observe trends, and understand possible relationships between lifestyle behavior and blood glucose fluctuation.",
    ]
    for item in english:
        p = doc.add_paragraph()
        set_para_format(p, first_indent=True)
        run = p.add_run(item)
        set_run_font(run, east="Times New Roman", west="Times New Roman", size=10.5)
    noindent_para(doc, "KEYWORDS: blood glucose management; Flutter; Supabase; health data analysis", bold=True)


def add_chapter_1(doc):
    heading(doc, "第一章 绪论", 1, page_break=True)
    heading(doc, "1.1 研究背景", 2)
    for t in [
        "血糖水平是反映人体能量代谢状态的重要指标之一。世界卫生组织在糖尿病事实资料中指出，糖尿病及高血糖相关疾病负担仍在增长，健康饮食、规律身体活动、保持适宜体重和避免烟草使用有助于预防或延缓 2 型糖尿病发生[1]。虽然本文系统面向普通人群而非糖尿病治疗场景，但这一背景说明，围绕血糖变化进行日常健康管理具有现实意义。",
        "在普通人的生活中，血糖并不是孤立变化的数值。高升糖指数食物可能引起餐后血糖较快升高，随后胰岛素分泌增加，部分个体可能出现明显回落、疲劳、饥饿或注意力下降等体验。相反，结构更均衡的饮食、适度运动和规律作息有助于降低波动幅度。将饮食、运动、主观状态与血糖记录关联起来，比单纯保存血糖数值更能帮助用户理解自己的生活习惯。",
        "移动应用为这种记录和反馈提供了较好的入口。智能手机具备随身携带、拍照、通知提醒、图表展示和云端同步等能力，适合承载轻量级健康管理场景。与专业医疗系统相比，面向普通用户的应用更强调低门槛、可解释、持续记录和隐私保护。本文选择血糖作为切入点，是因为它能够连接饮食、运动、精力状态和体重管理等多个维度，适合作为毕业设计中应用开发和数据分析结合的主题。",
        "需要强调的是，本文实现的系统不是医疗诊断工具，不提供疾病诊断、治疗方案或用药建议。系统中的参考区间、红黄绿灯提示和运动建议只用于生活方式观察，不能替代医生、营养师或专业医疗设备的判断。论文后续各章在需求分析、算法设计和结果讨论中均以这一边界为前提。",
    ]:
        para(doc, t)
    heading(doc, "1.2 国内外研究现状", 2)
    for t in [
        "目前健康管理应用大致可以分为三类：第一类侧重运动和体重管理，通过步数、运动时长和热量消耗估算用户活动量；第二类侧重饮食管理，通过食物库、图片识别或营养成分表估算能量摄入；第三类侧重慢性病指标记录，例如血压、血糖或用药提醒。多数应用已经具备良好的单项记录能力，但跨数据维度的解释能力仍然有限。",
        "糖尿病防治指南强调血糖管理需要综合饮食、运动、体重和自我监测等因素[2]。中国居民膳食指南也强调食物多样、合理搭配和控制添加糖摄入等原则[3]。这些公共健康资料为本系统的设计提供了方向：应用不应只提醒用户“少吃糖”，而应帮助用户在自己的记录中观察不同食物、份量和饭后活动对状态的影响。",
        "在饮食分析方面，血糖指数和血糖负荷研究为理解不同食物的餐后反应提供了重要参考。国际 GI/GL 表格研究系统整理了大量食物的升糖反应数据[4]，说明相同热量的食物也可能产生不同的血糖影响。然而，普通用户很难直接使用复杂表格，因此应用需要将这些概念转化为可理解的记录和提示，例如“继续观察”“减少频率”“可以保留”的红黄绿灯提示。",
        "在运动分析方面，MET 是估算身体活动能量消耗的常用指标。Compendium of Physical Activities 对不同活动类型给出了 MET 编码和参考值[5]。移动应用可以根据用户体重、运动时长和活动 MET 值估算能量消耗。本文系统使用 `exercise_catalog` 存储运动类型和 MET 值，保存运动记录时写入 `mets_snapshot`，使历史记录不受后续目录调整影响。",
    ]:
        para(doc, t)
    heading(doc, "1.3 研究意义", 2)
    for t in [
        "从应用价值看，本系统尝试把血糖记录从单点数据扩展为生活行为反馈。用户在饭后记录血糖、饮食和状态后，可以看到某类食物是否多次与血糖偏高、偏低或疲劳感相关，也可以获得饭后轻运动建议。这种反馈虽然不能用于诊断，但能够帮助用户形成更稳定的记录习惯，并逐步识别适合自己的饮食和运动模式。",
        "从工程实践看，本项目覆盖了一个移动健康应用从前端界面、状态管理、后端认证、数据库建模、对象存储、后端函数、第三方 API 调用、隐私保护到测试验证的完整链路。与仅使用本地存储的课程项目相比，Supabase BaaS 模式能够体现现代移动应用中“前端快速开发 + 云端安全能力”的特点。",
        "从毕业设计角度看，题目具有较强的综合性：Flutter 用于实现跨端界面，Provider 用于管理状态，Supabase 用于认证和数据存储，Edge Function 用于后端中转，AnalysisService 用于规则分析。论文可以围绕需求、架构、数据、安全、实现和测试展开，形成较完整的软件工程设计说明。",
    ]:
        para(doc, t)
    heading(doc, "1.4 论文主要工作", 2)
    for t in [
        "本文的主要工作包括以下几个方面。第一，分析普通用户在血糖健康管理中的记录、提醒、分析和隐私需求，明确系统不承担医疗诊断职责。第二，设计基于 Flutter 和 Supabase 的总体架构，说明客户端、状态管理、数据访问层、云端数据库和后端函数之间的职责划分。",
        "第三，设计用户档案、血糖、饮食、运动、状态、提醒和食物热量目录等数据模型，并通过 RLS 策略实现用户数据隔离。第四，实现饮食图片识别的安全中转流程，客户端只上传图片和调用 Edge Function，百度 API Key、Secret Key 和 Supabase service role key 均保存在服务端环境变量中。",
        "第五，实现基于餐后时间窗口和主观状态的红黄绿灯食物分析，以及基于 MET 的运动消耗估算。第六，通过静态分析、单元测试、组件测试和异常场景测试验证系统功能，最终形成可演示的应用界面和毕业论文初稿。",
    ]:
        para(doc, t)


def add_chapter_2(doc):
    heading(doc, "第二章 需求分析", 1, page_break=True)
    heading(doc, "2.1 用户需求与系统边界", 2)
    for t in [
        "本系统的目标用户是希望关注血糖波动、饮食结构、运动习惯和精力状态的普通人群。用户可能没有连续血糖监测设备，也不一定具备营养学或医学背景，因此系统必须降低记录门槛，避免复杂术语堆叠，并以清晰的文字、图表和颜色提示解释数据。系统面向毕业设计演示场景，要求功能闭环完整、数据结构清晰、安全边界可说明。",
        "系统边界包括两个外部参与者：普通用户和时间触发器/提醒服务。普通用户主动完成注册登录、记录数据、上传饮食图片、查看分析结果和维护个人资料；时间触发器/提醒服务则根据用户设置的本地通知时间提醒其记录饭后血糖或进行晚间回看。系统不会主动诊断疾病，也不会向第三方医疗机构发送用户数据。",
    ]:
        para(doc, t)
    add_picture(doc, "fig_use_case.png", "图 2-1 系统用例图", 5.9)
    heading(doc, "2.2 功能需求分析", 2)
    rows = [
        ["登录注册", "用户通过邮箱和密码注册、登录、退出；注册后创建默认档案、体重记录和提醒。"],
        ["首页概览", "展示最新血糖、记录数量、状态标签、红黄绿灯食物、饭后运动建议和近期记录。"],
        ["血糖记录", "支持手动录入血糖值、记录时间、时段、单位和数据来源，并查看或删除历史记录。"],
        ["饮食记录", "支持手动添加食物、拍照或相册识别、食物名称确认、克重和热量调整。"],
        ["运动记录", "从运动目录选择运动类型，输入时长，基于 MET 和体重估算消耗。"],
        ["状态记录", "记录疲劳、平稳、精力充沛等主观状态，可关联最近餐食或运动。"],
        ["分析建议", "基于餐后时间窗口、血糖状态和主观状态生成红黄绿灯食物提示。"],
        ["提醒设置", "支持新增、启停和删除本地每日提醒，帮助用户保持记录习惯。"],
    ]
    table(doc, "表 2-1 系统功能需求表", ["模块", "需求说明"], rows, widths=[3.2, 11.5])
    for t in [
        "血糖记录模块要求支持常见记录时段，例如空腹、早餐后、午餐前、午餐后、晚餐前、晚餐后、睡前、凌晨和随机。由于本系统主要服务普通用户和毕业设计演示，当前实现以手动录入为主，同时保留 `source = cgm` 的兼容字段，为未来接入连续血糖监测数据提供扩展空间。",
        "饮食记录模块是系统的重点之一。用户可以直接手动填写食物，也可以上传图片由后端调用百度菜品识别 API。识别结果只作为初始建议，用户仍需确认食物名称、选择常见份量或输入克重，并可手动覆盖最终热量。这种设计承认图片识别和典型热量存在误差，同时让用户以较低成本完成记录。",
        "分析模块不追求复杂医学模型，而是强调可解释规则。系统将餐食和餐后 30—180 分钟内的血糖记录进行配对，同时结合状态记录，统计食物多次出现后的血糖稳定性和主观体验。记录不足时提示继续观察，避免过早给出确定判断。",
    ]:
        para(doc, t)
    heading(doc, "2.3 非功能与安全需求", 2)
    for t in [
        "在易用性方面，系统需要尽量减少输入步骤。血糖记录通过滑块、时段选择和保存按钮完成；饮食记录允许用户先拍照识别，再在表单中修正；运动记录将复杂的 MET 计算隐藏在后台，只展示时长、运动类型和估算消耗。界面文案采用生活化表达，降低普通用户使用门槛。",
        "在可靠性方面，系统需要处理网络异常、后端表缺失、图片 bucket 未配置、百度 API 密钥缺失等错误。客户端不应直接抛出晦涩异常，而应转换为用户可以理解的提示。例如当 `meal-images` bucket 不存在时，页面提示食物图片存储未配置；当百度 secrets 缺失时，提示识别服务缺少百度 API 配置。",
        "在安全性方面，系统必须遵循最小权限原则。Flutter 客户端只注入 Supabase anon key，不保存 service role key 和百度密钥。用户数据通过 Supabase RLS 策略隔离，图片存储按用户 ID 目录划分，Edge Function 在服务端校验 Authorization 和 `storage_path` 前缀。截图和论文材料使用演示数据，姓名、头像、设备信息和健康数据均不使用真实个人信息。",
        "在可维护性方面，系统将业务计算放入 `AnalysisService`，数据访问集中在 `HealthRepository`，页面通过 `HealthProvider` 获取状态。这样的分层降低了 UI 与后端逻辑的耦合，也便于在测试中替换 Repository 或直接测试纯函数。",
    ]:
        para(doc, t)


def add_chapter_3(doc):
    heading(doc, "第三章 系统总体设计", 1, page_break=True)
    heading(doc, "3.1 技术路线与架构选择", 2)
    for t in [
        "本系统客户端采用 Flutter。Flutter 官方文档将其描述为支持 iOS、Android、Web 和桌面端代码复用的跨平台 UI 工具包，并采用分层架构和响应式界面模型[6]。本项目已经生成 Android、iOS、Web、Windows、macOS 和 Linux 平台目录，虽然毕业设计演示以移动端为主，但跨端能力为后续扩展提供了基础。",
        "状态管理选用 Provider。项目中 `HealthProvider` 负责保存最新用户档案、血糖历史、餐食历史、运动历史、状态历史、提醒列表和分析结果，并在数据刷新后通知页面重建。Provider 的职责是连接 UI 和 Repository，而不是直接拼接 SQL 或处理复杂业务计算。",
        "后端采用 Supabase BaaS。与传统自建 RESTful 后端相比，Supabase 提供 Auth、Postgres、Storage、Edge Function 和 RLS 等现成能力，客户端可以通过 `supabase_flutter` SDK 直接完成认证、数据库访问、对象存储上传和函数调用。supabase_flutter 是 Supabase 的 Flutter 客户端库，用于在 Flutter 应用中集成 Supabase 服务[11]。",
        "采用 BaaS 并不意味着放弃后端设计。相反，数据库 schema、RLS 策略、Storage bucket 策略和 Edge Function 的鉴权逻辑仍然需要开发者明确设计。本文系统的后端安全重点在于：普通 CRUD 通过 RLS 限制用户行级数据，饮食图片识别通过 Edge Function 中转第三方 API，所有敏感密钥只存在服务端环境变量或 Supabase Secrets。",
    ]:
        para(doc, t)
    heading(doc, "3.2 系统架构设计", 2)
    add_picture(doc, "fig_architecture.png", "图 3-1 系统总体架构图", 5.9)
    for t in [
        "系统从上到下可以划分为表示层、状态与业务层、BaaS 后端层和外部服务层。表示层由 `lib/pages` 下的页面组成，包括登录、首页、血糖、饮食、运动、分析、状态记录和个人设置。页面负责展示表单、列表、图表和按钮，不直接保存长期业务状态。",
        "状态与业务层由 Provider、Repository 和 Service 组成。`HealthProvider` 在登录后调用 `loadDashboardData()` 加载仪表盘数据，并将数据拆分为多个可供页面监听的列表。`HealthRepository` 封装 Supabase Auth、用户档案、血糖、饮食、运动、状态、提醒和图片识别相关 CRUD。`AnalysisService` 则封装热量计算、运动消耗、血糖状态判断、食物红黄绿灯和手动分析摘要。",
        "BaaS 后端层使用 Supabase。Auth 负责邮箱密码登录注册，Postgres 保存核心健康数据，Storage 的 `meal-images` 私有 bucket 保存饮食图片，Edge Function `recognize-meal` 调用百度菜品识别 API。外部服务层只包含百度菜品识别 API，不由 Flutter 直接访问。",
    ]:
        para(doc, t)
    heading(doc, "3.3 核心数据流设计", 2)
    add_picture(doc, "fig_data_flow.png", "图 3-2 核心数据流图", 5.9)
    for t in [
        "用户在页面录入数据后，页面调用 Provider 暴露的方法或直接使用 Provider 中的 Repository，Repository 再通过 supabase_flutter SDK 向后端提交数据。保存完成后，Provider 重新加载快照数据，页面根据新的状态展示首页指标、列表和分析建议。这一流程保证页面始终以 Repository 返回的数据为准，避免多处维护临时副本。",
        "分析流程不直接在数据库中执行，而是加载用户近期数据后在客户端本地运行。这样做的原因是当前分析规则较轻量，适合使用 Dart 纯函数实现和测试；同时规则中需要组合血糖、餐食和状态等多个列表，放在客户端可以快速迭代。若未来数据规模增大或需要跨用户统计，才需要迁移到数据库函数或后端任务。",
        "饮食图片识别流程是系统中唯一涉及外部 API 的链路。客户端先上传图片到 Supabase Storage，再调用 Edge Function。Edge Function 使用服务端密钥下载图片、获取百度 access_token 并调用识别接口。该流程使客户端仅接触用户会话和图片路径，不接触百度 API Key、Secret Key 或 Supabase service role key。",
    ]:
        para(doc, t)


def add_chapter_4(doc):
    heading(doc, "第四章 数据库与安全设计", 1, page_break=True)
    heading(doc, "4.1 数据模型设计", 2)
    for t in [
        "数据库设计围绕“用户—记录—分析”展开。`user_profile` 保存用户静态资料，`user_body_metrics` 保存体重变化，`blood_glucose_logs` 保存血糖记录，`meals` 和 `meal_items` 保存一次用餐及其食物明细，`exercise_catalog` 和 `exercise_logs` 保存运动类型与用户运动记录，`wellness_status` 保存主观状态，`reminder_settings` 保存提醒设置，`food_calorie_catalog` 保存常见食物热量目录。",
        "`meals` 与 `meal_items` 采用单据与明细结构。`meals.id` 是一次完整用餐的主键，包含用户 ID、餐次类型和用餐时间；`meal_items.meal_id` 指向 `meals.id`，保存识别原始名称、用户确认名称、每 100g 热量、克重、常见份量标签、用户覆盖热量和图片路径。`meal_items` 不冗余保存 `user_id`，权限通过关联的 `meals.user_id` 间接判断。",
        "`food_calorie_catalog` 是热量目录而不是强制外键表。系统查询该表是为了在用户输入食物名称后提供候选食物、别名、每 100g 热量和常见份量选项。最终保存到 `meal_items` 的仍是用户确认后的记录快照。这样做可以避免目录后续修改影响历史餐食，也方便用户记录不在目录中的食物。",
    ]:
        para(doc, t)
    add_picture(doc, "fig_er.png", "图 4-1 数据库简化 ER 图", 6.2)
    rows = [
        ["user_profile", "id", "用户档案", "与 auth.users 一一对应，保存昵称、性别、身高、出生日期。"],
        ["user_body_metrics", "user_id", "体重记录", "按时间记录体重，用于体重变化和运动消耗估算。"],
        ["blood_glucose_logs", "user_id", "血糖记录", "保存记录时间、时段、数值、单位和来源。"],
        ["meals", "user_id", "餐次记录", "保存一次用餐的餐次类型和用餐时间。"],
        ["meal_items", "meal_id", "餐食明细", "保存食物名称、热量、克重、用户覆盖值和图片路径。"],
        ["exercise_catalog", "id", "运动目录", "保存运动名称、MET 值、分类和描述，认证用户只读。"],
        ["exercise_logs", "user_id", "运动记录", "保存运动时间、运动类型、时长、mets_snapshot 和消耗。"],
        ["wellness_status", "user_id", "主观状态", "保存状态等级、备注，可关联最近餐食或运动。"],
        ["reminder_settings", "user_id", "提醒设置", "保存提醒时间、开关和标签。"],
        ["food_calorie_catalog", "id", "食物热量目录", "保存食物、别名、每 100g 热量和常见份量。"],
    ]
    table(doc, "表 4-1 核心数据表说明", ["数据表", "关键字段", "作用", "说明"], rows, widths=[3.2, 3.0, 3.0, 6.2])
    heading(doc, "4.2 后端安全与 RLS 策略", 2)
    for t in [
        "Supabase 官方文档指出，RLS 适用于需要细粒度授权规则的场景，并建议在暴露 schema 中启用 RLS[7]。本系统中所有用户私有表均启用 RLS，策略的基本思路是：直接包含 `user_id` 的表使用 `auth.uid() = user_id` 限制访问，`user_profile` 使用 `auth.uid() = id`，子表 `meal_items` 通过 `meals` 表间接校验所属用户。",
        "对于 `blood_glucose_logs`、`exercise_logs`、`wellness_status`、`reminder_settings` 等表，策略结构较一致。用户只能查询、插入、更新或删除自己的记录，不能通过客户端访问其他用户的数据。以下片段展示了血糖记录的典型策略写法：",
    ]:
        para(doc, t)
    code_block(doc, """
alter table public.blood_glucose_logs enable row level security;
create policy "glucose_own_rows" on public.blood_glucose_logs
  for all using (auth.uid() = user_id)
  with check (auth.uid() = user_id);
""")
    for t in [
        "`meal_items` 是 `meals` 的明细表，自身没有 `user_id`。如果简单为该表添加 `user_id`，会造成冗余和一致性风险；如果不做策略，又会暴露所有餐食明细。因此系统采用 `EXISTS` 子查询校验其父级餐次是否属于当前用户。策略片段如下：",
    ]:
        para(doc, t)
    code_block(doc, """
create policy "meal_items_own_rows" on public.meal_items
  for all using (
    exists (
      select 1 from public.meals
      where meals.id = meal_items.meal_id
        and meals.user_id = auth.uid()
    )
  )
  with check (
    exists (
      select 1 from public.meals
      where meals.id = meal_items.meal_id
        and meals.user_id = auth.uid()
    )
  );
""")
    for t in [
        "目录类表采用只读策略。`exercise_catalog` 和 `food_calorie_catalog` 对认证用户开放查询，但禁止普通客户端插入、更新和删除。这种设计既能让所有用户共用运动和食物基础数据，又能避免用户端随意修改公共目录影响他人记录。",
        "Storage 安全也依赖 RLS。Supabase Storage 文档说明，可以在 `storage.objects` 上创建策略并按 bucket 和文件夹限制访问[8]。本系统的 `meal-images` 是私有 bucket，图片路径以用户 ID 开头，上传和读取策略均要求 `auth.uid()::text = (storage.foldername(name))[1]`，从路径层面隔离不同用户图片。",
    ]:
        para(doc, t)
    heading(doc, "4.3 饮食图片存储与密钥管理", 2)
    for t in [
        "饮食图片识别涉及第三方 API，因此必须避免客户端直接持有百度密钥。Supabase Edge Functions 运行在 Deno 兼容环境，适合处理需要低延迟的认证 HTTP 端点和外部 API 编排任务[9]。Supabase 官方文档也明确 service role key 不应在浏览器中使用，而应放在服务端环境中；生产环境 secrets 可通过 Dashboard 或 CLI 设置[10]。",
        "本系统 Edge Function 使用的敏感变量包括 `SUPABASE_URL`、`SUPABASE_ANON_KEY`、`SUPABASE_SERVICE_ROLE_KEY`、`BAIDU_API_KEY` 和 `BAIDU_SECRET_KEY`。其中 `SUPABASE_ANON_KEY` 可以在启用 RLS 的前提下由客户端使用，`SUPABASE_SERVICE_ROLE_KEY` 只允许 Edge Function 使用，百度密钥只允许存放在 Supabase Secrets 中。",
        "在论文和代码片段中，所有密钥均以占位符展示，例如 `YOUR_SUPABASE_URL` 和 `YOUR_SUPABASE_ANON_KEY`。本地运行时通过 `--dart-define=SUPABASE_ANON_KEY=...` 注入客户端匿名 key，后端 secrets 通过 Supabase 管理，不写入 `.env`、`.vscode/launch.json` 或源码仓库。",
    ]:
        para(doc, t)


def add_chapter_5(doc):
    heading(doc, "第五章 功能模块实现", 1, page_break=True)
    heading(doc, "5.1 登录注册与个人档案", 2)
    for t in [
        "应用入口位于 `lib/main.dart`。程序启动时先初始化 Flutter 绑定、Supabase 配置和本地提醒服务，然后根据当前会话判断进入主页还是登录页。`GlucoseAssistantApp` 使用 Material 3 主题，并通过 `MultiProvider` 注入 `HealthProvider`，保证页面可以统一读取健康数据状态。",
        "登录注册模块使用 Supabase Auth 的邮箱和密码模式。用户注册时提交邮箱、密码和显示名称，Supabase Auth 创建用户后，数据库触发器 `public.handle_new_user()` 自动创建默认 `user_profile`、默认 `user_body_metrics` 和默认提醒。客户端的 `ensureCurrentUserProfile()` 作为兜底逻辑，在历史数据或断网恢复场景中补齐用户档案。",
        "个人资料页面允许用户维护显示名称、性别、身高、出生日期和体重。年龄不作为静态字段保存，而是由出生日期动态计算，避免每年手动更新。体重则作为时间序列保存到 `user_body_metrics`，便于未来扩展体重变化趋势，也为运动消耗计算提供最新体重。",
    ]:
        para(doc, t)
    heading(doc, "5.2 血糖、饮食、运动和状态记录", 2)
    for t in [
        "血糖记录页面提供滑块输入血糖值，并支持 `mmol/L` 和 `mg/dL` 两种单位字段。当前分析逻辑主要基于 `mmol/L`，但数据库保留单位字段以便未来转换。记录时段包括空腹、餐后、睡前、凌晨和随机等场景，数据来源包括 `manual` 和 `cgm`，其中 CGM 当前作为预留扩展。",
        "饮食记录页面支持手动添加多个食物，也支持从相机或相册选择图片。识别结果返回后，页面会显示原始识别名称、用户确认名称、每 100g 热量、克重、常见份量和最终热量。用户可以根据实际份量修正克重，也可以直接覆盖最终热量。保存时 Repository 先插入 `meals`，再批量插入 `meal_items`，如果明细保存失败则回滚已插入的餐次。",
        "运动记录页面从 `exercise_catalog` 读取运动类型和 MET 值。用户选择运动项目和时长后，系统查询最新体重并计算消耗热量。保存记录时同时写入运动 ID、时长、消耗热量和 `mets_snapshot`。这样即使未来修改运动目录中的 MET 参考值，历史记录仍能保留当时的计算依据。",
        "状态记录页面用于记录主观感受，包括极度疲劳、略感疲惫、状态平稳、感觉不错和精力充沛。状态可选关联最近一次餐食或运动，也可不关联。主观状态的价值在于补充血糖数值无法直接体现的体验信息，使分析模块可以观察食物、血糖和精力状态之间的组合关系。",
    ]:
        para(doc, t)
    heading(doc, "5.3 饮食识别与热量估算", 2)
    add_picture(doc, "fig_sequence.png", "图 5-1 饮食图片识别时序图", 6.2)
    for t in [
        "饮食图片识别流程分为五步。第一，用户在 Flutter 页面拍照或选择相册图片。第二，客户端将图片上传到 Supabase Storage 的 `meal-images` 私有 bucket，路径中包含当前用户 ID。第三，客户端调用 `recognize-meal` Edge Function，并传入 `storage_path`。第四，Edge Function 校验用户身份和路径前缀，使用 service role key 下载图片。第五，Edge Function 获取百度 access_token，并向百度菜品识别接口提交图片 Base64。",
        "百度菜品识别 API 要求通过 API Key 和 Secret Key 获取 access_token，并向 `https://aip.baidubce.com/rest/2.0/image-classify/v2/dish` 发送 POST 请求[12]。接口支持 `top_num` 和 `filter_threshold` 等参数，返回结果中包含菜名、热量和置信度字段。本系统将返回值转换为 `food_name_raw`、`calories_raw` 和 `confidence`，再交给客户端展示和修正。",
        "当前 Edge Function 按请求获取 access_token。百度示例文档也提示，线上环境中 access_token 有过期时间，客户端可自行缓存并在过期后重新获取[12]。本系统出于实现简洁性未做缓存，论文将其列为后续优化方向，以减少重复 token 请求带来的延迟。",
        "热量计算采用“每 100g 热量 + 实际克重 + 用户覆盖值”的方式。识别或食物库返回的 `calories_raw` 表示每 100g 热量，用户输入或选择的 `grams` 表示实际食物重量。当用户未覆盖最终热量时，数据库 generated column 使用公式：",
    ]:
        para(doc, t)
    formula(doc, "Calories_final = calories_raw × grams / 100")
    for t in [
        "当用户直接输入最终热量时，`calories_user_override` 优先于公式结果。这一设计承认识别和估算都可能存在误差，允许用户根据包装标识、外卖信息或个人判断修正结果。客户端保存时不写入 generated column，而是写入基础字段，由数据库自动生成最终热量，保证计算逻辑集中。",
    ]:
        para(doc, t)
    heading(doc, "5.4 分析建议与红黄绿灯规则", 2)
    for t in [
        "分析模块由 `AnalysisService` 实现，尽量保持纯函数形式，便于单元测试。血糖状态判断采用项目预设阈值：低于 3.9 mmol/L 视为偏低，高于 10.0 mmol/L 视为偏高，其他情况视为平稳。该阈值参考指南中常用血糖控制范围和项目演示需求，但系统不会据此诊断疾病。",
        "餐食与血糖的关联采用餐后时间窗口。系统遍历血糖记录，寻找记录时间在餐后 30—180 分钟内的最近餐次，并将该血糖状态计入对应餐食。状态记录则优先按 `related_meal_id` 直接关联，若无直接关联，则在 0—240 分钟窗口内寻找最近餐次。这样既能利用用户主动关联，也能减少普通用户忘记关联时的数据浪费。",
    ]:
        para(doc, t)
    rows = [
        ["黄灯", "无餐后配对血糖，或记录次数不足", "提示继续记录餐后血糖和状态，不做确定判断。"],
        ["红灯", "同一食物多次关联偏高/偏低血糖，或一次波动并伴随多次差状态", "建议减少频率、控制份量，并观察搭配或饭后活动。"],
        ["绿灯", "同一食物多次关联平稳血糖，且无差状态记录", "提示当前记录下较稳定，可以继续保留。"],
        ["黄灯", "既未达到红灯也未达到绿灯条件", "提示可以吃但继续观察，关注份量、时间和搭配。"],
    ]
    table(doc, "表 5-1 红黄绿灯规则阈值表", ["等级", "触发条件", "系统提示含义"], rows, widths=[2.0, 6.5, 6.4])
    for t in [
        "运动建议基于最新血糖和最近运动时长生成。如果最新血糖偏高，系统建议下一餐后先尝试 10—15 分钟慢走；如果最近两天运动记录偏少，则建议从饭后 12 分钟散步开始；如果餐食记录不足，则提示用户先补充饮食记录。该建议属于行为提醒，不构成医学处方。",
        "运动热量消耗采用 MET 公式：",
    ]:
        para(doc, t)
    formula(doc, "Calories = MET × weight(kg) × duration(h)")
    for t in [
        "在实现中，`duration` 以分钟录入，因此计算时使用 `duration_minutes / 60` 转换为小时。例如体重 65 kg、MET 为 4.3、运动 30 分钟时，估算消耗约为 140 kcal。系统保存记录时写入 `mets_snapshot`，保证历史记录可追溯。",
    ]:
        para(doc, t)
    heading(doc, "5.5 提醒与界面实现", 2)
    for t in [
        "提醒功能由 Supabase 中的 `reminder_settings` 表和 Flutter 本地通知共同实现。用户可以在个人设置页新增提醒时间、启停提醒或删除提醒。Provider 加载仪表盘数据后调用 `ReminderService.syncReminders()`，将云端提醒设置同步到本地通知计划。这样既能跨设备保留提醒配置，也能利用移动端本地通知能力。",
        "界面设计采用轻量卡片、指标 pill、环形指标和折线图展示核心信息。首页不采用复杂医学仪表盘，而是突出“最新血糖”“状态”“红黄绿灯食物”和“饭后轻运动”。血糖页强调快速保存，饮食页强调图片识别和人工确认，分析页强调原因解释和下一步建议。页面截图见图 5-2，截图数据均为演示数据并已脱敏。",
    ]:
        para(doc, t)
    add_picture(doc, "fig_screens.png", "图 5-2 系统主要页面实现效果", 6.4)


def add_chapter_6(doc):
    heading(doc, "第六章 测试与结果分析", 1, page_break=True)
    heading(doc, "6.1 测试环境与测试方法", 2)
    for t in [
        "本项目在 Windows 环境下使用 Flutter 工具链进行验证。验证命令包括 `flutter analyze` 和 `flutter test`。静态分析用于检查 Dart 代码中的类型、语法、lint 和潜在错误；测试命令用于执行单元测试、组件测试和部分异常逻辑测试。",
        "本次论文撰写前实际执行了验证命令。`flutter analyze` 输出 No issues found，说明当前项目在静态检查层面没有发现问题。`flutter test` 共执行 22 个测试，全部通过。测试结果可作为第六章功能正确性和代码质量的基础依据。",
        "由于当前运行环境没有配置 Supabase anon key，也不使用真实用户账号，论文中的界面截图采用临时演示数据生成，并在图注中说明。演示截图不包含真实姓名、头像、设备信息、联系方式或真实健康数据，避免论文材料泄露个人隐私。",
    ]:
        para(doc, t)
    rows = [
        ["开发框架", "Flutter / Dart", "跨端 UI、状态响应和页面构建。"],
        ["后端服务", "Supabase", "Auth、Postgres、Storage、Edge Function。"],
        ["状态管理", "provider", "统一管理用户健康快照和页面刷新。"],
        ["图表展示", "fl_chart", "首页血糖折线图。"],
        ["图片输入", "image_picker", "饮食拍照和相册选择。"],
        ["提醒", "flutter_local_notifications + timezone", "本地每日提醒。"],
        ["验证命令", "flutter analyze / flutter test", "静态检查和自动化测试。"],
    ]
    table(doc, "表 6-1 测试环境与依赖说明", ["项目", "工具或依赖", "用途"], rows, widths=[3.0, 5.2, 6.4])
    heading(doc, "6.2 功能、算法与安全测试", 2)
    rows = [
        ["静态分析", "执行 flutter analyze", "No issues found", "通过"],
        ["热量计算", "430 kcal/100g 与份量或克重计算", "得到预期热量", "通过"],
        ["用户覆盖热量", "存在 calories_user_override", "优先使用用户覆盖值", "通过"],
        ["MET 运动消耗", "MET=4.3、体重65kg、30分钟", "约 140 kcal", "通过"],
        ["红黄绿灯", "餐后窗口内多次偏高或平稳", "正确标记红/绿/黄", "通过"],
        ["血糖边界", "0.6 和 33.3 mmol/L 极端值", "不抛异常并参与分类", "通过"],
        ["提醒表缺失", "模拟 PGRST205", "提示 schema 缺失或返回空列表", "通过"],
        ["网络异常", "模拟 failed host lookup", "不误判为表缺失", "通过"],
        ["资料刷新", "模拟资料和体重返回", "Provider 正确更新", "通过"],
        ["按钮防连击", "连续点击保存运动", "只触发一次保存", "通过"],
    ]
    table(doc, "表 6-2 自动化测试用例摘要", ["测试类型", "输入或场景", "预期结果", "结果"], rows, widths=[3.0, 5.0, 5.2, 1.8])
    for t in [
        "算法测试主要覆盖三个风险点。第一是饮食热量计算。系统需要同时支持每 100g 热量、克重和用户覆盖最终热量，如果优先级错误，将直接影响首页和饮食历史展示。第二是运动消耗估算。公式虽然简单，但必须确保分钟到小时的单位换算正确。第三是红黄绿灯分类。该规则涉及餐食、血糖和状态三类数据，如果时间窗口或归因规则错误，用户看到的提示就会失真。",
        "容错测试覆盖 Supabase schema 缺失、网络异常和百度识别配置缺失等情况。健康管理应用不应在后端配置不完整时直接崩溃，而应给出明确提示，方便开发者和演示者定位问题。测试中对 `Bucket not found`、`Missing env: BAIDU_API_KEY` 和 `PGRST205` 等典型错误进行了友好提示验证。",
        "安全边界测试在当前自动化测试中以策略审查和异常模拟为主。论文中重点检查了 RLS 策略是否覆盖用户私有表、`meal_items` 是否通过父表校验用户、Storage bucket 是否按用户 ID 目录隔离，以及客户端是否没有写入 service role key 和百度密钥。后续若接入真实 Supabase 测试项目，可以进一步通过两个测试账号验证 A 用户无法读取 B 用户记录。",
        "跨端适配方面，当前项目已生成多个平台目录，UI 使用 SafeArea、ListView、Material 3 组件和响应式宽度约束，适合移动端纵向滚动展示。毕业设计演示主要建议在 Android 模拟器或真机完成，因为相机、相册和通知权限更接近实际使用场景。Web 和桌面端可作为后续适配方向继续检查。",
    ]:
        para(doc, t)
    heading(doc, "6.3 界面效果与结果分析", 2)
    for t in [
        "从界面结果看，系统将高频操作放在底部导航的一级页面中：首页、血糖、饮食、运动和分析。状态记录和个人设置虽然重要，但使用频率相对较低，因此通过首页入口进入。这样的结构符合普通用户“先看今日状态，再补充记录，再查看分析”的使用路径。",
        "从数据结果看，系统已经形成完整闭环。用户注册后有档案和提醒；用户记录血糖、饮食、运动和状态后，首页和分析页可以读取这些记录并生成反馈；用户删除记录后，Provider 会重新加载数据并刷新界面。该闭环说明系统不是单纯页面原型，而是具备可运行的数据流。",
        "从安全结果看，系统没有在 Flutter 源码中写入百度 API Key、Secret Key 或 Supabase service role key。Supabase anon key 通过运行参数注入，后端敏感变量由 Supabase Secrets 管理。饮食图片识别必须经过 Edge Function，且函数校验用户身份和图片路径前缀，降低了越权识别和密钥泄露风险。",
        "从隐私合规结果看，系统将真实用户数据、演示截图和论文材料进行了区分。数据库中的私有记录依赖 RLS 和用户 ID 隔离，论文截图只使用脱敏演示数据，测试输出中不打印密钥和用户完整身份信息。对于健康类应用而言，这种区分十分重要，因为即使系统不承担医疗诊断职责，血糖、体重、饮食和状态记录仍然属于敏感个人信息，开发和论文展示都应遵循最小披露原则。",
        "从局限看，当前分析规则仍然偏规则化和保守。红黄绿灯判断依赖用户连续记录，数据不足时只能提示继续观察；手动血糖记录无法呈现 CGM 那样的连续曲线；百度菜品识别返回的是典型热量，不代表用户当前盘中食物的精确营养分析。因此本文将系统定位为健康管理辅助工具，而不是医学判断系统。",
    ]:
        para(doc, t)


def add_chapter_7(doc):
    heading(doc, "第七章 结论与展望", 1, page_break=True)
    heading(doc, "7.1 结论", 2)
    for t in [
        "本文围绕“面向普通人群的血糖健康管理助手 APP 的设计与实现”完成了一套 Flutter + Supabase 的移动健康管理应用。系统覆盖登录注册、首页概览、血糖记录、饮食记录、图片识别、运动记录、状态记录、数据分析、个人设置和提醒等核心功能，能够支撑普通用户围绕血糖波动进行日常记录和趋势观察。",
        "在技术实现上，系统采用 Flutter 构建跨端界面，使用 Provider 管理页面状态，通过 HealthRepository 封装 Supabase 数据访问，通过 AnalysisService 实现可测试的业务计算。后端利用 Supabase Auth、Postgres、Storage、Edge Function 和 RLS 形成较完整的 BaaS 架构。饮食识别通过 Edge Function 安全中转百度菜品识别 API，避免客户端暴露敏感密钥。",
        "在数据设计上，系统将用户档案、体重、血糖、餐次、餐食明细、运动、主观状态和提醒拆分为多个关系清晰的数据表。`meals` 与 `meal_items` 采用单据明细结构，`exercise_logs` 保存 `mets_snapshot`，`meal_items` 保存识别和确认后的热量快照，这些设计使历史记录具备可追溯性。",
        "在分析设计上，系统采用餐后 30—180 分钟窗口、血糖状态和主观状态构建红黄绿灯食物提示，并基于 MET 公式估算运动消耗。该分析方式不追求医学诊断，而是强调可解释、可演示和可迭代。系统仅用于健康记录、趋势观察和生活方式辅助，不提供医疗诊断或治疗建议。",
        "在测试结果上，项目通过 `flutter analyze` 静态检查，`flutter test` 中 22 个测试全部通过，覆盖热量计算、运动消耗、红黄绿灯分类、异常提示、资料刷新和防连击等关键逻辑。结果表明，系统功能链路基本完整，具备毕业设计演示和论文撰写的工程基础。",
    ]:
        para(doc, t)
    heading(doc, "7.2 展望", 2)
    for t in [
        "第一，后续可以接入真实 CGM 数据。当前系统已经保留 `source = cgm` 字段和 CGM 分析视图，但尚未实现连续血糖设备接入。若未来接入传感器或第三方数据平台，可以分析血糖峰值、回落速度、日内稳定性和夜间低血糖风险，使分析粒度从手动关键点扩展到连续曲线。",
        "第二，可以优化饮食识别和热量估算。当前百度菜品识别返回的是典型热量，系统通过克重和用户覆盖值降低误差。后续可以引入更多食物库、条形码识别、包装营养成分识别或多食物分割，使热量估算更接近真实餐盘。",
        "第三，可以提升个性化分析能力。当前红黄绿灯规则以固定阈值和次数为主，适合毕业设计初版。未来可以在获得更多用户历史数据后，引入个体基线、餐食组合、运动间隔、睡眠状态和体重变化等因素，生成更细致的个性化建议。",
        "第四，可以加强后端性能和安全测试。当前 Edge Function 按请求获取百度 access_token，后续可增加服务端缓存与过期刷新机制；RLS 可通过测试项目中的双账号进行自动化验证；图片上传可增加 EXIF 清理、文件大小限制和内容类型校验，进一步降低隐私和滥用风险。",
        "第五，可以完善跨端发布能力。项目已经具备多平台目录，但正式发布前仍需更换 Android applicationId、配置 release 签名、补充 iOS 权限文案、完成相机和通知权限真机验证。对于毕业设计而言，当前版本已经能够展示系统设计与实现思路；对于真实产品而言，还需要持续完善合规、稳定性和用户体验。",
    ]:
        para(doc, t)


def add_references_appendix(doc):
    heading(doc, "参考文献", 1, page_break=True)
    refs = [
        "World Health Organization. Diabetes[EB/OL]. (2024-11-14)[2026-04-30]. https://www.who.int/news-room/fact-sheets/detail/diabetes.",
        "中华医学会糖尿病学分会. 中国2型糖尿病防治指南(2020年版)[J]. 中华糖尿病杂志, 2021, 13(4):315-409.",
        "中国营养学会. 中国居民膳食指南(2022)[M]. 北京: 人民卫生出版社, 2022.",
        "ATKINSON F S, BRAND-MILLER J C, FOSTER-POWELL K, et al. International tables of glycemic index and glycemic load values 2021: a systematic review[J]. The American Journal of Clinical Nutrition, 2021, 114(5):1625-1632.",
        "AINSWORTH B E, HASKELL W L, HERRMANN S D, et al. 2011 Compendium of Physical Activities: a second update of codes and MET values[J]. Medicine & Science in Sports & Exercise, 2011, 43(8):1575-1581.",
        "Flutter. Flutter architectural overview[EB/OL]. [2026-04-30]. https://docs.flutter.dev/resources/architectural-overview.",
        "Supabase. Row Level Security[EB/OL]. [2026-04-30]. https://supabase.com/docs/guides/database/postgres/row-level-security.",
        "Supabase. Storage Access Control[EB/OL]. [2026-04-30]. https://supabase.com/docs/guides/storage/security/access-control.",
        "Supabase. Edge Functions[EB/OL]. [2026-04-30]. https://supabase.com/docs/guides/functions.",
        "Supabase. Environment Variables[EB/OL]. [2026-04-30]. https://supabase.com/docs/guides/functions/secrets.",
        "Supabase. supabase_flutter package[EB/OL]. [2026-04-30]. https://pub.dev/packages/supabase_flutter.",
        "百度智能云. 菜品识别 API 文档[EB/OL]. [2026-04-30]. https://ai.baidu.com/ai-doc/IMAGERECOGNITION/tk3bcxbb0.",
    ]
    for i, ref in enumerate(refs, 1):
        p = doc.add_paragraph()
        set_para_format(p, first_indent=False, line=17, before=0, after=2)
        run = p.add_run(f"[{i}] {ref}")
        set_run_font(run, size=9)
    heading(doc, "附  录", 1, page_break=True)
    heading(doc, "附录 A 关键验证命令", 2)
    code_block(doc, """
flutter pub get
flutter analyze
flutter test
flutter run --dart-define=SUPABASE_ANON_KEY=YOUR_SUPABASE_ANON_KEY
""")
    para(doc, "上述命令中 `YOUR_SUPABASE_ANON_KEY` 为运行时占位符，不应写入源码仓库。涉及 Supabase service role key、百度 API Key 和 Secret Key 的配置均应通过 Supabase Secrets 管理。", collect=False)
    heading(doc, "附录 B 核心算法伪代码", 2)
    code_block(doc, """
for each glucose_record:
  meal = nearest meal where 30 <= glucose_time - meal_time <= 180 minutes
  if meal exists:
    attach glucose state to meal

for each food in meals:
  if glucose_count == 0:
    mark yellow
  else if unstable_count >= 2 or (unstable_count >= 1 and bad_status_count >= 2):
    mark red
  else if stable_count >= 2 and unstable_count == 0 and bad_status_count == 0:
    mark green
  else:
    mark yellow
""")
    heading(doc, "附录 C 演示数据与脱敏说明", 2)
    para(doc, "论文中的界面截图均为演示数据。截图中邮箱使用 demo@example.com，用户名称、头像、设备信息和健康数据均不对应真实个人；饮食、运动和血糖记录仅用于展示界面布局和系统流程。若后续使用真实设备截图，应先删除 EXIF 信息、遮挡账号信息，并确认没有真实健康数据泄露。", collect=False)
    heading(doc, "致  谢", 1, page_break=True)
    for t in [
        "在本课题完成过程中，感谢指导老师在选题方向、系统设计和论文结构方面给予的指导。老师对毕业设计的应用价值、技术实现完整性和论文规范性提出了要求，使我能够从单一功能开发转向更加完整的软件工程实现。",
        "感谢同学和朋友在应用体验、界面文字和演示流程方面提供反馈，帮助我发现了记录流程、异常提示和截图展示中的不足。感谢开源社区和官方文档提供的 Flutter、Supabase 等技术资料，使本项目能够在较短时间内完成跨端界面、云端数据和后端函数的整合。",
        "最后，感谢家人在毕业设计期间给予的支持。本文仍有不足之处，例如真实用户数据规模有限、CGM 接入尚未完成、分析规则仍较保守，后续我将继续改进系统的稳定性、个性化能力和实际应用价值。",
    ]:
        para(doc, t, collect=False)


def build_doc():
    generate_images()
    doc = Document()
    configure_document(doc)
    core = doc.core_properties
    core.author = ""
    core.last_modified_by = ""
    core.title = TITLE
    core.subject = "本科毕业论文初稿"
    core.keywords = "血糖管理; Flutter; Supabase; 健康数据分析"

    add_toc(doc)
    add_abstracts(doc)
    add_chapter_1(doc)
    add_chapter_2(doc)
    add_chapter_3(doc)
    add_chapter_4(doc)
    add_chapter_5(doc)
    add_chapter_6(doc)
    add_chapter_7(doc)
    add_references_appendix(doc)

    doc.save(OUT_DOCX)
    plain = "".join(body_text_parts)
    chinese_chars = len(re.findall(r"[\u4e00-\u9fff]", plain))
    print(f"Wrote {OUT_DOCX}")
    print(f"Approx Chinese body characters: {chinese_chars}")
    print(f"Image files: {len(list(IMG_DIR.glob('*.png')))}")


if __name__ == "__main__":
    build_doc()
